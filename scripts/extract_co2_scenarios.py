"""Extract annual global CO2 from the published Meinshausen et al. (2020) ZIP.

Usage: .venv/bin/python scripts/extract_co2_scenarios.py /path/to/supplement.zip
Download URL is recorded in data/co2/sources.json. No network access at runtime.
"""

import argparse
import hashlib
import io
import json
from pathlib import Path
import zipfile

import openpyxl
import pandas as pd


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("archive", type=Path)
    args = parser.parse_args()
    archive = args.archive.read_bytes()
    member = "SUPPLEMENT_DataTables_Meinshausen_6May2020.xlsx"
    with zipfile.ZipFile(io.BytesIO(archive)) as z:
        workbook = openpyxl.load_workbook(io.BytesIO(z.read(member)), read_only=True, data_only=True)

    def read_series(sheet, first, last):
        ws = workbook[sheet]
        if (ws["B9"].value, ws["B10"].value, ws["B11"].value) != ("CO2", "ppm", "World"):
            raise ValueError(f"Unexpected variable/unit/region in {sheet}")
        values = {int(year): float(co2) for year, co2 in ws.iter_rows(min_row=13, max_col=2, values_only=True)
                  if isinstance(year, (int, float)) and first <= year <= last}
        if sorted(values) != list(range(first, last + 1)):
            raise ValueError(f"Missing annual values in {sheet}")
        return values

    history_sheet = "T2 - History Year 1750 to 2014"
    history = read_series(history_sheet, 1900, 2014)
    sheets = {"SSP1-2.6": "T4 -  SSP1-2.6 ", "SSP2-4.5": "T5 - SSP2-4.5 ", "SSP5-8.5": "T11 - SSP5-8.5 "}
    df = pd.DataFrame({"year": range(1900, 2101)})
    for scenario, sheet in sheets.items():
        df[scenario] = df.year.map(history | read_series(sheet, 2015, 2100))
    output = Path(__file__).resolve().parents[1] / "data" / "co2"
    output.mkdir(parents=True, exist_ok=True)
    csv_path = output / "co2_concentrations_1900_2100.csv"
    df.to_csv(csv_path, index=False, float_format="%.9f")
    metadata = {
        "source_url": "https://gmd.copernicus.org/articles/13/3571/2020/gmd-13-3571-2020-supplement.zip",
        "source_archive_sha256": hashlib.sha256(archive).hexdigest(),
        "csv_sha256": hashlib.sha256(csv_path.read_bytes()).hexdigest(),
        "workbook": member, "historical_sheet": history_sheet, "scenario_sheets": sheets,
        "historical_doi": "10.5194/gmd-10-2057-2017", "scenario_doi": "10.5194/gmd-13-3571-2020",
        "variable": "CO2", "unit": "ppm", "region": "World", "frequency": "annual mean",
        "historical_years": [1900, 2014], "scenario_years": [2015, 2100],
        "license": "CC BY-SA 4.0 (dataset; see article Data availability)",
        "processing": "Extracted annual values, rounded to 9 decimal places in ppm. No fitting, bias adjustment or extrapolation. Scenario sheets describe 2015–2016 as observation based."
    }
    (output / "sources.json").write_text(json.dumps(metadata, indent=2, ensure_ascii=False) + "\n")
    print(df[df.year.isin([1900, 1980, 2000, 2014, 2050, 2080, 2100])].to_string(index=False))


if __name__ == "__main__":
    main()
